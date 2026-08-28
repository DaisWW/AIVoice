from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document

from app.script_parser import (
    ScriptFormatError,
    analyze_script_pronunciation,
    parse_content,
    parse_docx_sections,
    parse_guide,
    validate_docx_archive,
)


def test_parse_supported_text_and_csv_formats() -> None:
    text_items = parse_content(
        "台词甲 | mo-la？↗\n台词乙\tGU-la。↘\nmo……la……".encode("utf-8"),
        ".txt",
    )
    csv_items = parse_content(
        "text,pronunciation\n台词甲,mo-la\n".encode("utf-8"),
        ".csv",
    )

    assert [item.text for item in text_items[:2]] == ["台词甲", "台词乙"]
    assert text_items[0].direction == "rise"
    assert text_items[1].emphasis == ("gu",)
    assert text_items[2].text == "摸……啦……"
    assert csv_items[0].generated_text == "摸啦。"


@pytest.mark.parametrize("prefix", ["raw:", "ipa:", "phoneme:"])
def test_raw_pronunciation_keeps_custom_phonemes(prefix: str) -> None:
    item = parse_content(f"虫语角色 | {prefix} t͡ʃa-ʀ——ɬa↗".encode("utf-8"), ".txt")[0]

    assert item.raw_mode is True
    assert item.generated_text == "t͡ʃa ʀ ɬa?"
    assert item.hold_units == (0, 2, 0)
    assert item.emphasis == ()


@pytest.mark.parametrize(
    ("content", "suffix", "message"),
    [
        (b"text", ".exe", "仅支持"),
        (b"xx-unknown", ".txt", "没有拟声汉字映射"),
        (b"\xff\xfe\x00", ".txt", "UTF-8"),
        (b"not a docx", ".docx", "DOCX"),
    ],
)
def test_reject_invalid_scripts(content: bytes, suffix: str, message: str) -> None:
    with pytest.raises(ScriptFormatError, match=message):
        parse_content(content, suffix)


def test_parse_legacy_guide_groups_multiline_text(tmp_path: Path) -> None:
    guide = tmp_path / "guide.md"
    guide.write_text(
        "## 愚公\n\n【1 发音】mo-la\n第一行\n第二行\n\n表演：平静\n\n【静默】\n不生成\n",
        encoding="utf-8",
    )

    sections = parse_guide(guide)

    assert list(sections) == ["愚公"]
    assert sections["愚公"][0].text == "第一行\n第二行"
    assert sections["愚公"][0].pronunciation == "mo-la"


def test_parse_docx_sections_merges_adjacent_tables_by_role() -> None:
    document = Document()
    document.add_paragraph("角色甲")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "第一句"
    document.add_paragraph("")
    document.add_paragraph("角色甲")
    document.add_paragraph("")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "第二句"
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("角色乙")
    document.add_paragraph("")
    document.add_paragraph("第三句")

    output = io.BytesIO()
    document.save(output)
    sections = parse_docx_sections(output.getvalue())

    assert list(sections) == ["角色甲", "角色乙"]
    assert [item.text for item in sections["角色甲"]] == ["第一句", "第二句"]
    assert [item.text for item in sections["角色乙"]] == ["第三句"]


def test_parse_docx_sections_keeps_explicit_no_dialogue_roles() -> None:
    document = Document()
    document.add_paragraph("有台词角色")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "第一句"
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("无台词角色")
    document.add_paragraph("")
    document.add_paragraph("无台词")
    document.add_paragraph("（尖锐一点的鸟叫声）")

    output = io.BytesIO()
    document.save(output)
    sections = parse_docx_sections(output.getvalue())

    assert list(sections) == ["有台词角色", "无台词角色"]
    assert [item.text for item in sections["无台词角色"]] == ["尖锐一点的鸟叫声"]


def test_parse_docx_sections_keeps_inline_no_dialogue_description() -> None:
    document = Document()
    document.add_paragraph("无台词角色")
    document.add_paragraph("")
    document.add_paragraph("无台词（只能哼两声）")

    output = io.BytesIO()
    document.save(output)
    sections = parse_docx_sections(output.getvalue())

    assert [item.text for item in sections["无台词角色"]] == ["只能哼两声"]


@pytest.mark.parametrize("member_name", ["../word/document.xml", "/word/document.xml"])
def test_docx_archive_rejects_unsafe_member_paths(member_name: str) -> None:
    content = io.BytesIO()
    with zipfile.ZipFile(content, "w") as archive:
        archive.writestr(member_name, "x")

    with pytest.raises(ScriptFormatError, match="压缩包条目无效"):
        validate_docx_archive(content.getvalue())


@pytest.mark.parametrize(
    ("pronunciation", "expected"),
    [
        ("shu—la", (1, 0)),
        ("shu——la", (2, 0)),
        ("shu–la", (1, 0)),
        ("shu―la", (1, 0)),
        ("shu－la", (1, 0)),
        ("嗯——……ha", (2, 0)),
    ],
)
def test_long_dashes_encode_hold_units(
    pronunciation: str, expected: tuple[int, ...]
) -> None:
    item = parse_content(pronunciation.encode("utf-8"), ".txt")[0]

    assert item.hold_units == expected
    assert not any(dash in item.generated_text for dash in "—–―－")


def test_prosody_indices_include_direct_chinese_syllables() -> None:
    analysis = analyze_script_pronunciation("嗯——……GU-la")

    assert analysis.emphasis_indices == (1,)
    assert analysis.hold_units == (2, 0, 0)


@pytest.mark.parametrize("pronunciation", ["—mo", "mo———————"])
def test_invalid_hold_marks_are_rejected(pronunciation: str) -> None:
    with pytest.raises(ScriptFormatError, match="延音"):
        parse_content(pronunciation.encode("utf-8"), ".txt")
